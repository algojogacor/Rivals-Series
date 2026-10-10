import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client
import pymupdf

sys.stdout.reconfigure(encoding='utf-8')

# Palette derived directly from the provided cover artwork:
# - Background: Deep Midnight Tactical Navy (#0C1728)
# - Crimson Accent Bar: (#E53430)
# - Bone White Text: (#EEE6DA)
# - Slate Wireframe Lines: (#64748B / #7D8794)
# - Charcoal Body: (#1E293B)
# - Surface Tint: (#F8FAFC)
# - Subtle Rules: (#E2E8F0)

HEX_DARK_PRIMARY = "0C1728"      # Deep Midnight Tactical Navy
HEX_CRIMSON = "E53430"           # Competition Crimson Accent
HEX_DARK_SECONDARY = "334155"    # Slate 700 / Dark Slate Subhead
HEX_DARK_BODY = "1E293B"         # Slate 800 / Charcoal Body Text
HEX_MUTED = "64748B"             # Slate 500 / Muted Metadata & Headers
HEX_BG_TINT = "F8FAFC"           # Slate 50 / Subtle Shading Tint
HEX_BORDER_RULE = "CBD5E1"       # Slate 300 / Editorial Divider Hairline
HEX_INNER_RULE = "E2E8F0"        # Slate 200 / Subtle Interior Table Rule
HEX_WHITE = "FFFFFF"

COLOR_DARK_PRIMARY = RGBColor(0x0C, 0x17, 0x28)
COLOR_CRIMSON = RGBColor(0xE5, 0x34, 0x30)
COLOR_DARK_SECONDARY = RGBColor(0x33, 0x41, 0x55)
COLOR_DARK_BODY = RGBColor(0x1E, 0x29, 0x3B)
COLOR_MUTED = RGBColor(0x64, 0x74, 0x8B)

FONT_FAMILY = "Arial"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=90, bottom=90, left=110, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, val in borders.items():
        if val:
            sz, b_type, col = val
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{b_type}" w:sz="{sz}" w:space="0" w:color="{col}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def add_inline_runs(paragraph, text, base_size=Pt(10.5), base_color=COLOR_DARK_BODY, base_font=FONT_FAMILY):
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = base_font
            run.font.size = base_size
            run.font.bold = True
            run.font.color.rgb = COLOR_DARK_PRIMARY
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = base_font
            run.font.size = base_size
            run.font.italic = True
            run.font.color.rgb = base_color
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size.pt * 0.95)
            run.font.bold = True
            run.font.color.rgb = COLOR_DARK_PRIMARY
        else:
            run = paragraph.add_run(token)
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color

def format_booktabs_table(doc, col_widths, headers, rows_data, align_right_cols=[], align_center_cols=[]):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="{HEX_DARK_PRIMARY}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_DARK_PRIMARY}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="2" w:space="0" w:color="{HEX_INNER_RULE}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)
    
    # Header Row
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.width = Inches(col_widths[c_idx])
        set_cell_background(cell, HEX_BG_TINT)
        set_cell_margins(cell, top=100, bottom=100, left=110, right=110)
        set_cell_borders(
            cell,
            top=(8, "single", HEX_DARK_PRIMARY),
            bottom=(6, "single", HEX_CRIMSON),
            left=None,
            right=None
        )
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
        if c_idx in align_right_cols:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif c_idx in align_center_cols:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
        r = p.add_run(h_text)
        r.font.name = FONT_FAMILY
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK_PRIMARY
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data Rows
    for r_idx, r_data in enumerate(rows_data):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        bg_col = HEX_BG_TINT if (r_idx % 2 == 1) else HEX_WHITE
        is_last_row = (r_idx == len(rows_data) - 1)
        bottom_border = (8, "single", HEX_DARK_PRIMARY) if is_last_row else None
        
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=110, right=110)
            set_cell_borders(cell, top=None, bottom=bottom_border, left=None, right=None)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            if c_idx in align_right_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif c_idx in align_center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            add_inline_runs(p, str(val), base_size=Pt(9.0), base_color=COLOR_DARK_BODY)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def format_technical_form_box(doc, title, lines_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.378)
    set_cell_background(cell, HEX_BG_TINT)
    set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
    set_cell_borders(
        cell,
        left=(16, "single", HEX_DARK_PRIMARY),
        top=(4, "single", HEX_BORDER_RULE),
        bottom=(4, "single", HEX_BORDER_RULE),
        right=(4, "single", HEX_BORDER_RULE)
    )
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(4)
    r_pip = p0.add_run("■ ")
    r_pip.font.name = FONT_FAMILY
    r_pip.font.size = Pt(8.5)
    r_pip.font.color.rgb = COLOR_CRIMSON
    
    r0 = p0.add_run(title)
    r0.font.name = FONT_FAMILY
    r0.font.size = Pt(10)
    r0.font.bold = True
    r0.font.color.rgb = COLOR_DARK_PRIMARY
    
    for line in lines_text:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.25
        
        if ":" in line and not line.startswith("="):
            parts = line.split(":", 1)
            r_k = p.add_run(parts[0] + ":")
            r_k.font.name = "Consolas"
            r_k.font.size = Pt(9.0)
            r_k.font.bold = True
            r_k.font.color.rgb = COLOR_DARK_PRIMARY
            
            r_v = p.add_run(parts[1])
            r_v.font.name = "Consolas"
            r_v.font.size = Pt(9.0)
            r_v.font.color.rgb = COLOR_DARK_BODY
        else:
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9.0)
            if line.startswith("==") or line.startswith("--"):
                r.font.color.rgb = COLOR_MUTED
            elif line.startswith("[") and line.endswith("]"):
                r.font.bold = True
                r.font.color.rgb = COLOR_DARK_PRIMARY
            else:
                r.font.color.rgb = COLOR_DARK_BODY

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def format_checklist_box(doc, title, items):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.378)
    set_cell_background(cell, HEX_BG_TINT)
    set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
    set_cell_borders(
        cell,
        left=(16, "single", HEX_CRIMSON),
        top=(4, "single", HEX_BORDER_RULE),
        bottom=(4, "single", HEX_BORDER_RULE),
        right=(4, "single", HEX_BORDER_RULE)
    )
    
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_before = Pt(0)
    p0.paragraph_format.space_after = Pt(4)
    r_pip = p0.add_run("■ ")
    r_pip.font.name = FONT_FAMILY
    r_pip.font.size = Pt(8.5)
    r_pip.font.color.rgb = COLOR_CRIMSON
    
    r0 = p0.add_run(title)
    r0.font.name = FONT_FAMILY
    r0.font.size = Pt(10)
    r0.font.bold = True
    r0.font.color.rgb = COLOR_DARK_PRIMARY
    
    for chk_title, chk_subitems in items:
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(3.0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.25
        
        r_box = p.add_run("[  ] ")
        r_box.font.name = "Consolas"
        r_box.font.size = Pt(9.5)
        r_box.font.bold = True
        r_box.font.color.rgb = COLOR_CRIMSON
        
        r_t = p.add_run(chk_title)
        r_t.font.name = FONT_FAMILY
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = COLOR_DARK_PRIMARY
        
        for sub in chk_subitems:
            ps = cell.add_paragraph()
            ps.paragraph_format.left_indent = Inches(0.3)
            ps.paragraph_format.space_before = Pt(1.0)
            ps.paragraph_format.space_after = Pt(1.5)
            ps.paragraph_format.line_spacing = 1.25
            
            r_s = ps.add_run(sub)
            r_s.font.name = FONT_FAMILY
            r_s.font.size = Pt(9.0)
            r_s.font.color.rgb = COLOR_DARK_BODY

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_header_footer_redesigned(section):
    section.different_first_page_header_footer = False
    
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs:
        p.text = ""
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
    htbl = header.add_table(rows=1, cols=2, width=Inches(6.378))
    htbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    htbl.autofit = False
    
    tblPr_h = htbl._tbl.tblPr
    tblPr_h.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="3" w:color="{HEX_INNER_RULE}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    ))
    
    hc0 = htbl.cell(0, 0)
    hc0.width = Inches(4.2)
    set_cell_margins(hc0, top=0, bottom=30, left=0, right=0)
    hp0 = hc0.paragraphs[0]
    hp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp0.paragraph_format.space_before = Pt(0)
    hp0.paragraph_format.space_after = Pt(0)
    
    rdot = hp0.add_run("■ ")
    rdot.font.name = FONT_FAMILY
    rdot.font.size = Pt(7.0)
    rdot.font.color.rgb = COLOR_CRIMSON
    
    hr0 = hp0.add_run("RIVALS SERIES · SEASON 1")
    hr0.font.name = FONT_FAMILY
    hr0.font.size = Pt(8.5)
    hr0.font.bold = True
    hr0.font.color.rgb = COLOR_DARK_PRIMARY
    
    hc1 = htbl.cell(0, 1)
    hc1.width = Inches(2.178)
    set_cell_margins(hc1, top=0, bottom=30, left=0, right=0)
    hp1 = hc1.paragraphs[0]
    hp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp1.paragraph_format.space_before = Pt(0)
    hp1.paragraph_format.space_after = Pt(0)
    hr1 = hp1.add_run("OFFICIAL REGULATIONS")
    hr1.font.name = FONT_FAMILY
    hr1.font.size = Pt(8.5)
    hr1.font.color.rgb = COLOR_MUTED

    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        p.text = ""
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        
    ftbl = footer.add_table(rows=1, cols=2, width=Inches(6.378))
    ftbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    ftbl.autofit = False
    
    tblPr_f = ftbl._tbl.tblPr
    tblPr_f.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="3" w:color="{HEX_INNER_RULE}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    ))
    
    fc0 = ftbl.cell(0, 0)
    fc0.width = Inches(4.2)
    set_cell_margins(fc0, top=30, bottom=0, left=0, right=0)
    fp0 = fc0.paragraphs[0]
    fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp0.paragraph_format.space_before = Pt(0)
    fp0.paragraph_format.space_after = Pt(0)
    fr0 = fp0.add_run("REGULASI RESMI KOMPETISI · VERSI 1.0-DRAFT")
    fr0.font.name = FONT_FAMILY
    fr0.font.size = Pt(8.5)
    fr0.font.color.rgb = COLOR_MUTED
    
    fc1 = ftbl.cell(0, 1)
    fc1.width = Inches(2.178)
    set_cell_margins(fc1, top=30, bottom=0, left=0, right=0)
    fp1 = fc1.paragraphs[0]
    fp1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp1.paragraph_format.space_before = Pt(0)
    fp1.paragraph_format.space_after = Pt(0)
    fr1 = fp1.add_run("Halaman ")
    fr1.font.name = FONT_FAMILY
    fr1.font.size = Pt(8.5)
    fr1.font.color.rgb = COLOR_MUTED
    
    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    fp1._p.append(fld1)
    
    fr_mid = fp1.add_run(" dari ")
    fr_mid.font.name = FONT_FAMILY
    fr_mid.font.size = Pt(8.5)
    fr_mid.font.color.rgb = COLOR_MUTED
    
    fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"/>')
    fp1._p.append(fld2)

def build_cover_page_redesigned(doc, cover_img_path):
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(0)
    sec.bottom_margin = Mm(0)
    sec.left_margin = Mm(0)
    sec.right_margin = Mm(0)
    sec.header_distance = Mm(0)
    sec.footer_distance = Mm(0)
    
    sec.header.is_linked_to_previous = False
    for p in sec.header.paragraphs: p.text = ""
    sec.footer.is_linked_to_previous = False
    for p in sec.footer.paragraphs: p.text = ""
    
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    
    run = p.add_run()
    run.font.size = Pt(1)
    run.add_picture(cover_img_path, width=Mm(210), height=Mm(297))
    
    # Convert inline shape to floating anchor positioned relative to page (0, 0)
    # This guarantees ZERO blank white space at top, left, bottom, right!
    drawing = run._r.xpath('.//w:drawing')[0]
    inline = drawing.xpath('.//wp:inline')[0]
    
    extent = inline.xpath('.//wp:extent')[0]
    docPr = inline.xpath('.//wp:docPr')[0]
    cNvGraphicFramePr = inline.xpath('.//wp:cNvGraphicFramePr')[0]
    graphic = inline.xpath('.//a:graphic')[0]
    
    anchor_xml = (
        f'<wp:anchor {nsdecls("wp")} {nsdecls("a")} distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'</wp:anchor>'
    )
    anchor = parse_xml(anchor_xml)
    anchor.append(extent)
    anchor.append(parse_xml(f'<wp:wrapNone {nsdecls("wp")}/>'))
    anchor.append(docPr)
    anchor.append(cNvGraphicFramePr)
    anchor.append(graphic)
    
    drawing.remove(inline)
    drawing.append(anchor)

def build_doc_control_page_redesigned(doc):
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(8)
    p_h1.paragraph_format.space_after = Pt(2)
    p_h1.paragraph_format.keep_with_next = True
    
    r_pip = p_h1.add_run("■ ")
    r_pip.font.name = FONT_FAMILY
    r_pip.font.size = Pt(11)
    r_pip.font.color.rgb = COLOR_CRIMSON
    
    r = p_h1.add_run("KONTROL DOKUMEN DAN STATUS REGULASI")
    r.font.name = FONT_FAMILY
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr = p_h1._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="8" w:color="{HEX_DARK_PRIMARY}"/></w:pBdr>'))
    
    headers = ["Parameter", "Ketentuan Administratif"]
    rows = [
        ["Nama kompetisi", "Rivals Series Season 1 (RS S1)"],
        ["Status dokumen", "Draft konsolidasi resmi untuk persetujuan kedua peserta"],
        ["Versi dokumen", "1.0-draft"],
        ["Format pertandingan", "eFootball™ Dream Team · Friend Match 1v1"],
        ["Jumlah pertandingan reguler", "20 Matchday (MD1–MD20)"],
        ["Pihak peserta", "Peserta A dan Peserta B"],
        ["Competition Administrator", "Ditetapkan bersama oleh kedua peserta"],
        ["Official Website", "Akan ditambahkan setelah website tersedia"],
        ["Tanggal berlaku", "Setelah kedua peserta menyetujui dokumen ini"],
        ["Sumber aturan primer", "rules.md (Konsolidasi mandiri kedua peserta)"],
    ]
    format_booktabs_table(doc, col_widths=[2.1, 4.28], headers=headers, rows_data=rows)
    
    p_pr_h = doc.add_paragraph()
    p_pr_h.paragraph_format.space_before = Pt(14)
    p_pr_h.paragraph_format.space_after = Pt(2)
    p_pr_h.paragraph_format.keep_with_next = True
    r_pr_h = p_pr_h.add_run("Prinsip Utama Sumber Aturan")
    r_pr_h.font.name = FONT_FAMILY
    r_pr_h.font.size = Pt(10.5)
    r_pr_h.font.bold = True
    r_pr_h.font.color.rgb = COLOR_DARK_PRIMARY
    
    p_bq = doc.add_paragraph()
    p_bq.paragraph_format.left_indent = Inches(0.25)
    p_bq.paragraph_format.space_before = Pt(4)
    p_bq.paragraph_format.space_after = Pt(8)
    p_bq.paragraph_format.line_spacing = 1.35
    pPr_bq = p_bq._p.get_or_add_pPr()
    pPr_bq.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="16" w:space="8" w:color="{HEX_CRIMSON}"/></w:pBdr>'))
    
    r_bq = p_bq.add_run("Dokumen ini merupakan sumber aturan utama RS S1 berdasarkan rules.md. Buku resmi, website, pengumuman, dan lampiran harus konsisten dengan rules.md. Tidak ada pihak yang boleh mengubah aturan, hasil, atau sanksi secara sepihak.")
    r_bq.font.name = FONT_FAMILY
    r_bq.font.size = Pt(10)
    r_bq.font.italic = True
    r_bq.font.color.rgb = COLOR_DARK_BODY

def build_toc_page_redesigned(doc, toc_data):
    doc.add_page_break()
    
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(8)
    p_h1.paragraph_format.space_after = Pt(2)
    p_h1.paragraph_format.keep_with_next = True
    
    r_pip = p_h1.add_run("■ ")
    r_pip.font.name = FONT_FAMILY
    r_pip.font.size = Pt(11)
    r_pip.font.color.rgb = COLOR_CRIMSON
    
    r = p_h1.add_run("DAFTAR ISI")
    r.font.name = FONT_FAMILY
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr = p_h1._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="8" w:color="{HEX_DARK_PRIMARY}"/></w:pBdr>'))
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    
    for code, title, p_num in toc_data:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.15
        
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.378), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        
        r_code = p.add_run(f"{code}  —  ")
        r_code.font.name = FONT_FAMILY
        r_code.font.size = Pt(9.5)
        r_code.font.bold = True
        r_code.font.color.rgb = COLOR_DARK_PRIMARY
        
        r_title = p.add_run(title)
        r_title.font.name = FONT_FAMILY
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = COLOR_DARK_BODY
        
        r_tab = p.add_run("\t")
        
        r_page = p.add_run(f"Halaman {p_num}")
        r_page.font.name = FONT_FAMILY
        r_page.font.size = Pt(9.0)
        r_page.font.bold = True
        r_page.font.color.rgb = COLOR_DARK_SECONDARY

def build_approval_page_redesigned(doc):
    doc.add_page_break()
    
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(8)
    p_h1.paragraph_format.space_after = Pt(2)
    p_h1.paragraph_format.keep_with_next = True
    
    r_pip = p_h1.add_run("■ ")
    r_pip.font.name = FONT_FAMILY
    r_pip.font.size = Pt(11)
    r_pip.font.color.rgb = COLOR_CRIMSON
    
    r = p_h1.add_run("HALAMAN PERSETUJUAN")
    r.font.name = FONT_FAMILY
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr = p_h1._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="8" w:color="{HEX_DARK_PRIMARY}"/></w:pBdr>'))
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(6)
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.paragraph_format.line_spacing = 1.5
    r_intro = p_intro.add_run("Dengan menyetujui regulasi ini, kedua peserta menyatakan akan mengikuti RS S1 berdasarkan ketentuan yang tercantum dalam dokumen ini.")
    r_intro.font.name = FONT_FAMILY
    r_intro.font.size = Pt(10.5)
    r_intro.font.color.rgb = COLOR_DARK_BODY
    
    signatories = [
        ("PESERTA A", "Nama: ______________________________\nPersetujuan / Tanda Tangan: ______________________________\nTanggal: ______________________________"),
        ("PESERTA B", "Nama: ______________________________\nPersetujuan / Tanda Tangan: ______________________________\nTanggal: ______________________________"),
        ("COMPETITION ADMINISTRATOR", "Nama: ______________________________\nTanggal Penetapan: ______________________________"),
    ]
    
    for idx, (title, lines_text) in enumerate(signatories):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        
        row = tbl.rows[0]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        cell = row.cells[0]
        cell.width = Inches(6.378)
        set_cell_background(cell, HEX_BG_TINT)
        set_cell_margins(cell, top=100, bottom=100, left=130, right=130)
        set_cell_borders(
            cell,
            left=(14, "single", HEX_DARK_SECONDARY),
            top=(4, "single", HEX_BORDER_RULE),
            bottom=(4, "single", HEX_BORDER_RULE),
            right=(4, "single", HEX_BORDER_RULE)
        )
        
        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(4)
        r0 = p0.add_run(title)
        r0.font.name = FONT_FAMILY
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_DARK_PRIMARY
        
        for l in lines_text.split('\n'):
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.3
            r = p.add_run(l)
            r.font.name = FONT_FAMILY
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK_BODY
            
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)

def build_appendices_redesigned(doc):
    doc.add_page_break()
    
    # 0. MAIN APPENDICES BANNER
    p_main = doc.add_paragraph()
    p_main.paragraph_format.space_before = Pt(10)
    p_main.paragraph_format.space_after = Pt(2)
    p_main.paragraph_format.keep_with_next = True
    
    r_pip0 = p_main.add_run("■ ")
    r_pip0.font.name = FONT_FAMILY
    r_pip0.font.size = Pt(11)
    r_pip0.font.color.rgb = COLOR_CRIMSON
    
    r_m = p_main.add_run("LAMPIRAN OPERASIONAL KOMPETISI (APPENDIX A–E)")
    r_m.font.name = FONT_FAMILY
    r_m.font.size = Pt(13.5)
    r_m.font.bold = True
    r_m.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr_m = p_main._p.get_or_add_pPr()
    pPr_m.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="8" w:color="{HEX_DARK_PRIMARY}"/></w:pBdr>'))
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(6)
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_intro,
        "Bagian ini memuat format operasional resmi **Appendix A sampai Appendix E** untuk **Rivals Series Season 1 (RS S1)**. Seluruh template, jadwal laga, panduan formula klasemen, formulir pendaftaran sengketa, log amendemen, dan checklist pra-kompetisi tercantum secara lengkap dan terpadu di bawah ini sebagai instrumen administratif yang mengikat kedua belah pihak.",
        base_size=Pt(10.5)
    )

    # APPENDIX A
    p_app_a = doc.add_paragraph()
    p_app_a.paragraph_format.space_before = Pt(14)
    p_app_a.paragraph_format.space_after = Pt(2)
    p_app_a.paragraph_format.keep_with_next = True
    
    r_pip_a = p_app_a.add_run("■ ")
    r_pip_a.font.name = FONT_FAMILY
    r_pip_a.font.size = Pt(10)
    r_pip_a.font.color.rgb = COLOR_CRIMSON
    
    r_a = p_app_a.add_run("LAMPIRAN A — OFFICIAL FIXTURE LIST (MD1–MD20)")
    r_a.font.name = FONT_FAMILY
    r_a.font.size = Pt(12)
    r_a.font.bold = True
    r_a.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr_a = p_app_a._p.get_or_add_pPr()
    pPr_a.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="6" w:color="{HEX_BORDER_RULE}"/></w:pBdr>'))
    
    p_desc_a = doc.add_paragraph()
    p_desc_a.paragraph_format.space_before = Pt(4)
    p_desc_a.paragraph_format.space_after = Pt(8)
    p_desc_a.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_desc_a,
        "**Ketentuan Operasional Jadwal:** Urutan pertandingan bersifat fleksibel dan dapat dimainkan tidak berurutan sesuai kesepakatan kedua peserta. Pengaturan laga (*match settings*) wajib mengikuti ketentuan nomor matchday terkait (MD1–MD10: Group Stage, 8 Menit, ET OFF, PK OFF; MD11–MD20: Knockout Stage, 8 Menit, ET ON, PK ON).",
        base_size=Pt(10.0)
    )
    
    headers_a = ["MD", "Fase", "Jadwal", "Status", "Skor", "Catatan"]
    rows_a = [
        ["MD1", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD2", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD3", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD4", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD5", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD6", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD7", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD8", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD9", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD10", "Group Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD11", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD12", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD13", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD14", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD15", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD16", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD17", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD18", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD19", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
        ["MD20", "Knockout Stage", "Fleksibel", "Scheduled", "—", "—"],
    ]
    format_booktabs_table(
        doc,
        col_widths=[0.8, 1.4, 1.0, 1.1, 0.7, 1.378],
        headers=headers_a,
        rows_data=rows_a,
        align_center_cols=[0, 2, 3, 4]
    )
    
    p_note_a = doc.add_paragraph()
    p_note_a.paragraph_format.space_before = Pt(4)
    p_note_a.paragraph_format.space_after = Pt(12)
    p_note_a.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_note_a,
        "**Jendela Perubahan Skuad Pertengahan Musim (Pasal 15):** Di antara MD10 dan MD11, kedua peserta berhak melakukan pembaruan registrasi skuad (maksimal 26 kartu pemain) sebelum kick-off MD11 dimulai.",
        base_size=Pt(10.0)
    )

    # APPENDIX B
    doc.add_page_break()
    p_app_b = doc.add_paragraph()
    p_app_b.paragraph_format.space_before = Pt(8)
    p_app_b.paragraph_format.space_after = Pt(2)
    p_app_b.paragraph_format.keep_with_next = True
    
    r_pip_b = p_app_b.add_run("■ ")
    r_pip_b.font.name = FONT_FAMILY
    r_pip_b.font.size = Pt(10)
    r_pip_b.font.color.rgb = COLOR_CRIMSON
    
    r_b = p_app_b.add_run("LAMPIRAN B — OFFICIAL MATCH REPORT")
    r_b.font.name = FONT_FAMILY
    r_b.font.size = Pt(12)
    r_b.font.bold = True
    r_b.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr_b = p_app_b._p.get_or_add_pPr()
    pPr_b.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="6" w:color="{HEX_BORDER_RULE}"/></w:pBdr>'))
    
    p_desc_b = doc.add_paragraph()
    p_desc_b.paragraph_format.space_before = Pt(4)
    p_desc_b.paragraph_format.space_after = Pt(8)
    p_desc_b.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_desc_b,
        "Format laporan resmi hasil pertandingan yang wajib dikirimkan ke grup koordinasi resmi selambat-lambatnya 24 jam setelah pertandingan berakhir:",
        base_size=Pt(10.0)
    )
    
    report_lines = [
        "==================================================",
        "⚽ RIVALS SERIES S1 — OFFICIAL MATCH REPORT",
        "==================================================",
        "Matchday         : MD[..]",
        "Fase             : [Group Stage / Knockout Stage]",
        "Tanggal Main     : [DD/MM/YYYY, Pukul ..:.. WIB]",
        "Peserta A        : [Nama Peserta A]",
        "Peserta B        : [Nama Peserta B]",
        "",
        "SKOR AKHIR       : [Peserta A] ... - ... [Peserta B]",
        "Hasil            : [Peserta A Menang / Seri / Peserta B Menang / Forfeit / Technical Win]",
        "Status           : [Menunggu Konfirmasi / Disahkan / Disengketakan]",
        "",
        "Kredit Indonesia :",
        "- Peserta A      : [+1 kredit babak ke-.. / +2 kredit / 0 kredit]",
        "- Peserta B      : [+1 kredit babak ke-.. / +2 kredit / 0 kredit]",
        "",
        "Gangguan/Rekonstruksi : [Tidak Ada / Ada — jelaskan ringkas]",
        "Catatan Administrasi  : [Catatan tambahan jika ada]",
        "",
        "KONFIRMASI HASIL :",
        "- Peserta A      : [Ya / Belum]",
        "- Peserta B      : [Ya / Belum]",
        "==================================================",
        "📎 LAMPIRAN WAJIB:",
        "1. Screenshot layar hasil akhir pertandingan.",
        "2. Rekaman kandidat Goal of the Season (bila ada, simpan dalam 24 jam).",
        "--------------------------------------------------",
        "ℹ️ PENGINGAT BUKTI SUSUNAN PEMAIN (Pasal 20):",
        "Screenshot Starting XI & DSP wajib disimpan sebelum kick-off.",
        "Audit manual tidak diwajibkan; screenshot dibuka bila ada sengketa.",
        "==================================================",
    ]
    format_technical_form_box(doc, "Template Match Report Resmi (Konsol Teks / WhatsApp)", report_lines)
    
    p_rules_b = doc.add_paragraph()
    p_rules_b.paragraph_format.space_before = Pt(4)
    p_rules_b.paragraph_format.space_after = Pt(12)
    p_rules_b.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_rules_b,
        "**Ketentuan Bukti Susunan Pemain (Pasal 20):** Kedua peserta wajib mengambil tangkapan layar susunan pemain dan daftar susunan pemain (DSP) sebelum kick-off. Tidak ada kewajiban audit manual rutin setiap laga; bukti disimpan dan hanya diperiksa apabila terjadi pengajuan sengketa resmi oleh pihak lawan.",
        base_size=Pt(10.0)
    )

    # APPENDIX C
    doc.add_page_break()
    p_app_c = doc.add_paragraph()
    p_app_c.paragraph_format.space_before = Pt(8)
    p_app_c.paragraph_format.space_after = Pt(2)
    p_app_c.paragraph_format.keep_with_next = True
    
    r_pip_c = p_app_c.add_run("■ ")
    r_pip_c.font.name = FONT_FAMILY
    r_pip_c.font.size = Pt(10)
    r_pip_c.font.color.rgb = COLOR_CRIMSON
    
    r_c = p_app_c.add_run("LAMPIRAN C — OFFICIAL LEAGUE TABLE & FORMULA KLASEMEN")
    r_c.font.name = FONT_FAMILY
    r_c.font.size = Pt(12)
    r_c.font.bold = True
    r_c.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr_c = p_app_c._p.get_or_add_pPr()
    pPr_c.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="6" w:color="{HEX_BORDER_RULE}"/></w:pBdr>'))
    
    p_desc_c = doc.add_paragraph()
    p_desc_c.paragraph_format.space_before = Pt(4)
    p_desc_c.paragraph_format.space_after = Pt(6)
    p_desc_c.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_desc_c,
        "**Website Resmi:** `[URL ditambahkan setelah website tersedia]` (Pasal 48). Klasemen resmi dipublikasikan secara terbuka dan memuat 10 parameter wajib sebagai berikut:",
        base_size=Pt(10.0)
    )
    
    headers_c = ["Parameter", "Simbol", "Ketentuan Operasional & Perhitungan"]
    rows_c = [
        ["Matches Played", "MP", "Total laga yang telah dimainkan dan disahkan (Maksimal 20 laga)"],
        ["Wins", "W", "Jumlah kemenangan (Bernilai 3 poin)"],
        ["Draws", "D", "Jumlah hasil imbang (Bernilai 1 poin; hanya berlaku di MD1–MD10)"],
        ["Losses", "L", "Jumlah kekalahan (Bernilai 0 poin)"],
        ["Goals For", "GF", "Total seluruh gol memasukkan selama waktu reguler dan perpanjangan waktu (ET)"],
        ["Goals Against", "GA", "Total seluruh gol kemasukan selama waktu reguler dan perpanjangan waktu (ET)"],
        ["Goal Difference", "GD", "Selisih gol resmi dengan formula: GD = GF − GA"],
        ["Total Points", "PTS", "Akumulasi poin: Menang = 3 poin; Seri = 1 poin; Kalah = 0 poin"],
        ["Kredit Indonesia", "ID Credits", "Progres kredit babak pemain Indonesia per peserta dari target wajib 5 kredit"],
        ["Catatan Sanksi", "Sanctions", "Pengurangan poin aktif atau catatan administratif hasil sengketa"],
    ]
    format_booktabs_table(
        doc,
        col_widths=[1.5, 0.9, 3.978],
        headers=headers_c,
        rows_data=rows_c,
        align_center_cols=[1]
    )
    
    p_rules_c = doc.add_paragraph()
    p_rules_c.paragraph_format.space_before = Pt(4)
    p_rules_c.paragraph_format.space_after = Pt(12)
    p_rules_c.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_rules_c,
        "**Ketentuan Khusus Penghitungan Poin dan Skor:**\n"
        "1. **Poin Fase Knockout (MD11–MD20):** Wajib menghasilkan pemenang melalui ET dan adu penalti; pemenang meraih 3 poin dan pihak kalah 0 poin.\n"
        "2. **Pengecualian Adu Penalti:** Gol pada babak adu penalti **tidak dimasukkan** ke dalam perhitungan GF maupun GA (Pasal 9 Ayat 2).\n"
        "3. **Skor Forfeit Administratif:** Forfeit dicatat dengan skor administratif **0–3** (masuk ke GF/GA tim, tanpa pencatatan gol individu).\n"
        "4. **Technical Win (Pasal 28 Ayat 4):** Menang 3 poin tanpa menciptakan gol fiktif (GF dan GA tetap 0).\n"
        "5. **Sanksi Kuota Pemain Indonesia (Pasal 19 Ayat 5):** Peserta yang tidak memenuhi target minimal 5 kredit babak sepanjang musim dikenai sanksi pengurangan 3 poin di klasemen akhir.",
        base_size=Pt(10.0)
    )

    # APPENDIX D
    doc.add_page_break()
    p_app_d = doc.add_paragraph()
    p_app_d.paragraph_format.space_before = Pt(8)
    p_app_d.paragraph_format.space_after = Pt(2)
    p_app_d.paragraph_format.keep_with_next = True
    
    r_pip_d = p_app_d.add_run("■ ")
    r_pip_d.font.name = FONT_FAMILY
    r_pip_d.font.size = Pt(10)
    r_pip_d.font.color.rgb = COLOR_CRIMSON
    
    r_d = p_app_d.add_run("LAMPIRAN D — DISCIPLINARY AND DISPUTE REGISTER")
    r_d.font.name = FONT_FAMILY
    r_d.font.size = Pt(12)
    r_d.font.bold = True
    r_d.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr_d = p_app_d._p.get_or_add_pPr()
    pPr_d.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="6" w:color="{HEX_BORDER_RULE}"/></w:pBdr>'))
    
    p_desc_d = doc.add_paragraph()
    p_desc_d.paragraph_format.space_before = Pt(4)
    p_desc_d.paragraph_format.space_after = Pt(8)
    p_desc_d.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_desc_d,
        "Format pelaporan dan pencatatan sengketa resmi untuk dinilai secara internal atau diteruskan ke AI Adjudicator (BAB IX, Pasal 33–38):",
        base_size=Pt(10.0)
    )
    
    dispute_lines = [
        "==================================================",
        "⚖️ RIVALS SERIES S1 — DISPUTE / DISCIPLINARY REGISTER",
        "==================================================",
        "Nomor Perkara       : DISP-S1-[Nomor Urut]",
        "Matchday Terkait    : MD[..]",
        "Pihak Pengaju       : [Peserta A / Peserta B]",
        "Pihak Terlapor      : [Peserta B / Peserta A]",
        "Tanggal Pengajuan   : [DD/MM/YYYY, Pukul ..:.. WIB]",
        "Batas Respons (48j) : [DD/MM/YYYY, Pukul ..:.. WIB]",
        "",
        "PASAL YANG DIAJUKAN :",
        "[Contoh: Pasal 17 (Kartu Trending) / Pasal 18 (DSP ID) / Pasal 32 (Birdcaging) / dsb.]",
        "",
        "RINGKASAN DUGAAN    :",
        "[Uraian singkat fakta objektif yang dipersoalkan]",
        "",
        "BUKTI PENDUKUNG (Dilampirkan):",
        "1. [Screenshot susunan pemain pra-kick-off / Screenshot skor / dsb.]",
        "2. [Rekaman video gameplay kontinu min 30 detik untuk birdcaging]",
        "3. [Bukti tangkapan layar chat / log teknis terkait]",
        "",
        "TANGGAPAN PIHAK TERLAPOR:",
        "[Ringkasan tanggapan dan bukti tandingan dalam batas 48 jam]",
        "",
        "PROSES AI ADJUDICATOR (Bila sengketa dilanjutkan):",
        "- Layanan/Model AI  : [Nama model yang telah disepakati sebelum MD1]",
        "- Tanggal Adjudikasi: [DD/MM/YYYY]",
        "",
        "KESIMPULAN PUTUSAN  : [Terbukti / Tidak Terbukti / Bukti Tidak Memadai]",
        "PASAL DITERAPKAN    : [Sebutkan pasal regulasi yang menjadi dasar putusan]",
        "KONSEKUENSI/SANKSI  : [Status Quo / Peringatan / -1 Poin / Forfeit 0-3 / Technical Win]",
        "STATUS PERKARA      : [Terbuka / Menunggu Respons / Final Ditutup]",
        "LOG ADMINISTRASI    : [Catatan eksekusi perubahan klasemen di website]",
        "==================================================",
    ]
    format_technical_form_box(doc, "Formulir Register Sengketa Resmi (DISP-S1)", dispute_lines)
    
    p_rules_d = doc.add_paragraph()
    p_rules_d.paragraph_format.space_before = Pt(4)
    p_rules_d.paragraph_format.space_after = Pt(12)
    p_rules_d.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_rules_d,
        "**Prinsip Beban Pembuktian dan Status Quo (Pasal 36):** Beban pembuktian berada sepenuhnya pada pihak pengaju. Laporan wajib diajukan maksimal 24 jam setelah laga berakhir, dan pihak terlapor memiliki waktu 48 jam untuk memberikan tanggapan resmi. Jika bukti tidak memadai atau tidak memenuhi unsur objektif pasal, tuduhan ditolak dan **status quo tetap berlaku mutlak** (hasil yang telah disahkan tidak berubah).",
        base_size=Pt(10.0)
    )

    # APPENDIX E
    doc.add_page_break()
    p_app_e = doc.add_paragraph()
    p_app_e.paragraph_format.space_before = Pt(8)
    p_app_e.paragraph_format.space_after = Pt(2)
    p_app_e.paragraph_format.keep_with_next = True
    
    r_pip_e = p_app_e.add_run("■ ")
    r_pip_e.font.name = FONT_FAMILY
    r_pip_e.font.size = Pt(10)
    r_pip_e.font.color.rgb = COLOR_CRIMSON
    
    r_e = p_app_e.add_run("LAMPIRAN E — REGULATIONS AMENDMENT LOG & PRE-MD1 CHECKLIST")
    r_e.font.name = FONT_FAMILY
    r_e.font.size = Pt(12)
    r_e.font.bold = True
    r_e.font.color.rgb = COLOR_DARK_PRIMARY
    
    pPr_e = p_app_e._p.get_or_add_pPr()
    pPr_e.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="6" w:color="{HEX_BORDER_RULE}"/></w:pBdr>'))
    
    p_sub_e1 = doc.add_paragraph()
    p_sub_e1.paragraph_format.space_before = Pt(6)
    p_sub_e1.paragraph_format.space_after = Pt(2)
    p_sub_e1.paragraph_format.keep_with_next = True
    r_sub_e1 = p_sub_e1.add_run("1. Log Amendemen Regulasi (Amendment Log)")
    r_sub_e1.font.name = FONT_FAMILY
    r_sub_e1.font.size = Pt(11)
    r_sub_e1.font.bold = True
    r_sub_e1.font.color.rgb = COLOR_DARK_SECONDARY
    
    p_desc_e1 = doc.add_paragraph()
    p_desc_e1.paragraph_format.space_before = Pt(2)
    p_desc_e1.paragraph_format.space_after = Pt(6)
    p_desc_e1.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_desc_e1,
        "Sesuai Pasal 51, setiap perubahan regulasi wajib disetujui secara tertulis oleh kedua belah peserta dan dicatat dalam log berikut:",
        base_size=Pt(10.0)
    )
    
    headers_e = ["Versi", "Tanggal", "Pasal", "Uraian Perubahan", "Alasan Teknis", "Persetujuan A", "Persetujuan B", "Berlaku Mulai"]
    rows_e = [
        ["1.0-draft", "—", "—", "Draft Konsolidasi Resmi Awal", "Penyusunan regulasi awal", "Menunggu", "Menunggu", "Setelah disetujui"],
        ["[1.1]", "[DD/MM/YY]", "[Pasal ..]", "[Deskripsi klausul yang diubah]", "[Alasan perubahan]", "[Setuju/Tgl]", "[Setuju/Tgl]", "[Mulai MD..]"],
        ["[1.2]", "[DD/MM/YY]", "[Pasal ..]", "[Deskripsi klausul yang diubah]", "[Alasan perubahan]", "[Setuju/Tgl]", "[Setuju/Tgl]", "[Mulai MD..]"],
    ]
    format_booktabs_table(
        doc,
        col_widths=[0.7, 0.75, 0.6, 1.35, 1.15, 0.65, 0.65, 0.528],
        headers=headers_e,
        rows_data=rows_e,
        align_center_cols=[0, 1, 2, 5, 6, 7]
    )
    
    p_sub_e2 = doc.add_paragraph()
    p_sub_e2.paragraph_format.space_before = Pt(10)
    p_sub_e2.paragraph_format.space_after = Pt(2)
    p_sub_e2.paragraph_format.keep_with_next = True
    r_sub_e2 = p_sub_e2.add_run("2. Checklist Penetapan Wajib Sebelum Kick-off MD1")
    r_sub_e2.font.name = FONT_FAMILY
    r_sub_e2.font.size = Pt(11)
    r_sub_e2.font.bold = True
    r_sub_e2.font.color.rgb = COLOR_DARK_SECONDARY
    
    p_desc_e2 = doc.add_paragraph()
    p_desc_e2.paragraph_format.space_before = Pt(2)
    p_desc_e2.paragraph_format.space_after = Pt(6)
    p_desc_e2.paragraph_format.line_spacing = 1.5
    add_inline_runs(
        p_desc_e2,
        "Kedua peserta dan administrator wajib menyepakati dan melengkapi 6 instrumen pra-kompetisi berikut sebelum pertandingan MD1 dimulai:",
        base_size=Pt(10.0)
    )
    
    checklist_items = [
        ("1. Identitas Resmi Peserta (Pasal 1 & 11):", [
            "• Peserta A: Nama resmi, nama in-game eFootball™, dan nomor kontak WhatsApp telah terdaftar.",
            "• Peserta B: Nama resmi, nama in-game eFootball™, dan nomor kontak WhatsApp telah terdaftar."
        ]),
        ("2. Competition Administrator (Pasal 3 & 49):", [
            "• Penetapan figur Competition Administrator bersama yang disepakati oleh kedua peserta."
        ]),
        ("3. Model / Layanan AI Adjudicator (Pasal 38):", [
            "• Pemilihan model AI resmi (contoh: Claude 3.5 Sonnet / GPT-4o) untuk sengketa tak terselesaikan.",
            "• Penyimpanan template prompt adjudikasi standar yang identik oleh kedua pihak."
        ]),
        ("4. Alat Undian Koin Digital (Pasal 28):", [
            "• Pemilihan aplikasi/web koin digital dan penguncian sisi permanen: Sisi HEADS = Peserta A | Sisi TAILS = Peserta B."
        ]),
        ("5. Registrasi Skuad 26 Kartu (Pasal 14):", [
            "• Penguncian skuad maksimal 26 kartu Peserta A dan Peserta B sebelum kick-off MD1 (tanpa kartu kembar base_player_id).",
            "• Verifikasi minimal 2 kartu bertipe Trending dan pemain berkebangsaan Indonesia pada skuad terdaftar."
        ]),
        ("6. Pengesahan Dokumen Regulasi (BAB XII):", [
            "• Penandatanganan Halaman Persetujuan oleh Peserta A, Peserta B, dan Competition Administrator."
        ]),
    ]
    format_checklist_box(doc, "Checklist Administrasi Pra-Kompetisi (Pre-MD1 Verification)", checklist_items)

def parse_and_render_body_redesigned(doc, body_lines):
    idx = 0
    total = len(body_lines)
    is_first_bab = True
    current_pasal_num = 0
    in_pasal_definitions = False
    
    while idx < total:
        line = body_lines[idx].rstrip('\r\n')
        s = line.strip()
        
        if not s:
            idx += 1
            continue
            
        # 1. Heading 1: BAB
        if s.startswith('# BAB '):
            if is_first_bab:
                doc.add_page_break()
                is_first_bab = False
            
            parts = s.lstrip('#').strip().split('—')
            bab_label = parts[0].strip()
            bab_title = parts[1].strip() if len(parts) > 1 else ""
            
            p_bab = doc.add_paragraph()
            p_bab.paragraph_format.space_before = Pt(24)
            p_bab.paragraph_format.space_after = Pt(2)
            p_bab.paragraph_format.keep_with_next = True
            
            r_pip = p_bab.add_run("■ ")
            r_pip.font.name = FONT_FAMILY
            r_pip.font.size = Pt(9.5)
            r_pip.font.color.rgb = COLOR_CRIMSON
            
            r_bl = p_bab.add_run(bab_label)
            r_bl.font.name = FONT_FAMILY
            r_bl.font.size = Pt(11)
            r_bl.font.bold = True
            r_bl.font.color.rgb = COLOR_MUTED
            
            p_bt = doc.add_paragraph()
            p_bt.paragraph_format.space_before = Pt(0)
            p_bt.paragraph_format.space_after = Pt(12)
            p_bt.paragraph_format.keep_with_next = True
            
            r_bt = p_bt.add_run(bab_title)
            r_bt.font.name = FONT_FAMILY
            r_bt.font.size = Pt(13.5)
            r_bt.font.bold = True
            r_bt.font.color.rgb = COLOR_DARK_PRIMARY
            
            pPr_bt = p_bt._p.get_or_add_pPr()
            pPr_bt.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="6" w:color="{HEX_BORDER_RULE}"/></w:pBdr>'))
            
            idx += 1
            continue
            
        # 2. Heading 2: Pasal
        if s.startswith('## Pasal '):
            m_p = re.match(r'^## Pasal (\d+)', s)
            if m_p:
                current_pasal_num = int(m_p.group(1))
                in_pasal_definitions = (current_pasal_num == 2)
            else:
                in_pasal_definitions = False
                
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            
            pasal_text = s.lstrip('#').strip()
            run = p.add_run(pasal_text)
            run.font.name = FONT_FAMILY
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_DARK_PRIMARY
            idx += 1
            continue
            
        # 3. Heading 3: Ayat / Sub-bagian
        if s.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            
            ayat_text = s.lstrip('#').strip()
            run = p.add_run(ayat_text)
            run.font.name = FONT_FAMILY
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_DARK_SECONDARY
            idx += 1
            continue
            
        # 4. Divider '---'
        if s.startswith('---'):
            idx += 1
            continue
            
        # 5. Markdown Table Parser
        if s.startswith('|') and s.endswith('|'):
            table_lines = []
            while idx < total and body_lines[idx].strip().startswith('|') and body_lines[idx].strip().endswith('|'):
                table_lines.append(body_lines[idx].strip())
                idx += 1
                
            if len(table_lines) >= 3:
                raw_headers = [c.strip() for c in table_lines[0].strip('|').split('|')]
                raw_rows = []
                for tl in table_lines[2:]:
                    cols = [c.strip() for c in tl.strip('|').split('|')]
                    raw_rows.append(cols)
                    
                num_cols = len(raw_headers)
                total_w = 6.378
                if num_cols == 2:
                    c_widths = [3.0, 3.378]
                elif num_cols == 3:
                    c_widths = [2.0, 2.0, 2.378]
                elif num_cols == 4:
                    c_widths = [1.5, 1.5, 1.5, 1.878]
                else:
                    c_widths = [total_w / num_cols] * num_cols
                    
                align_r = []
                if num_cols == 2 and any("kredit" in h.lower() or "poin" in h.lower() for h in raw_headers):
                    align_r = [1]
                    c_widths = [5.0, 1.378]
                    
                format_booktabs_table(doc, c_widths, raw_headers, raw_rows, align_right_cols=align_r)
            continue

        # 6. Blockquote '> '
        if s.startswith('>'):
            bq_text = s.lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.35
            pPr = p._p.get_or_add_pPr()
            pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="16" w:space="8" w:color="{HEX_CRIMSON}"/></w:pBdr>'))
            add_inline_runs(p, bq_text, base_size=Pt(10), base_color=COLOR_DARK_BODY)
            idx += 1
            continue

        # 7. Numbered Legal Provision: '1. ', '2. ', etc.
        m_num = re.match(r'^(\d+)\.\s+(.*)$', s)
        if m_num:
            num_str = m_num.group(1)
            content = m_num.group(2)
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2.0)
            p.paragraph_format.space_after = Pt(6.0)
            p.paragraph_format.line_spacing = 1.5
            
            if in_pasal_definitions:
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.first_line_indent = Inches(-0.3)
                run_num = p.add_run(f"{num_str}. ")
                run_num.font.name = FONT_FAMILY
                run_num.font.size = Pt(10.5)
                run_num.font.bold = True
                run_num.font.color.rgb = COLOR_DARK_PRIMARY
            else:
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.35)
                run_num = p.add_run(f"({num_str}) ")
                run_num.font.name = FONT_FAMILY
                run_num.font.size = Pt(10.5)
                run_num.font.bold = True
                run_num.font.color.rgb = COLOR_DARK_PRIMARY
                
            add_inline_runs(p, content, base_size=Pt(10.5))
            idx += 1
            continue

        # 8. Lettered Subparagraph: 'a. ', 'b. ', etc.
        m_alpha = re.match(r'^([a-z])\.\s+(.*)$', s)
        if m_alpha:
            alpha_str = m_alpha.group(1)
            content = m_alpha.group(2)
            
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.55)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(5.0)
            p.paragraph_format.line_spacing = 1.5
            
            run_lbl = p.add_run(f"{alpha_str}. ")
            run_lbl.font.name = FONT_FAMILY
            run_lbl.font.size = Pt(10.5)
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = COLOR_DARK_SECONDARY
            
            add_inline_runs(p, content, base_size=Pt(10.5))
            idx += 1
            continue

        # 9. Bullet list: '- ' or '* '
        if s.startswith('- ') or s.startswith('* '):
            content = s[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(5.0)
            p.paragraph_format.line_spacing = 1.5
            
            run_lbl = p.add_run("• ")
            run_lbl.font.name = FONT_FAMILY
            run_lbl.font.size = Pt(10.5)
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = COLOR_DARK_SECONDARY
            
            add_inline_runs(p, content, base_size=Pt(10.5))
            idx += 1
            continue

        # 10. Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2.0)
        p.paragraph_format.space_after = Pt(6.0)
        p.paragraph_format.line_spacing = 1.5
        add_inline_runs(p, s, base_size=Pt(10.5))
        idx += 1

def generate_document(docx_path, pdf_path, cover_img_path, toc_data):
    print("Reading rules.md...")
    with open('rules.md', 'r', encoding='utf-8') as f:
        full_text = f.read()
        
    lines = full_text.splitlines()
    bab1_idx = -1
    approval_idx = -1
    
    for i, l in enumerate(lines):
        if l.strip().startswith('# BAB I '):
            bab1_idx = i
            break
            
    for i, l in enumerate(lines):
        if l.strip().startswith('# HALAMAN PERSETUJUAN'):
            approval_idx = i
            break
            
    body_lines = lines[bab1_idx:approval_idx]
    
    doc = docx.Document()
    
    # 1. Build Cover Page (Section 1 with floating zero-margin anchor)
    print("Building Cover Page with zero-margin floating anchor...")
    build_cover_page_redesigned(doc, cover_img_path)
    
    # 2. Add Section Break for Body (Section 2)
    sec_body = doc.add_section(WD_SECTION.NEW_PAGE)
    sec_body.page_width = Mm(210)
    sec_body.page_height = Mm(297)
    sec_body.top_margin = Mm(24)
    sec_body.bottom_margin = Mm(24)
    sec_body.left_margin = Mm(24)
    sec_body.right_margin = Mm(24)
    
    add_header_footer_redesigned(sec_body)
    
    # 3. Build Document Control & Dedicated TOC
    print("Building Document Control & Dedicated TOC...")
    build_doc_control_page_redesigned(doc)
    build_toc_page_redesigned(doc, toc_data)
    
    # 4. Parse & Render BAB I through BAB XII with 1.5 line spacing
    print("Parsing & Rendering BAB I through BAB XII with 1.5 line spacing...")
    parse_and_render_body_redesigned(doc, body_lines)
    
    # 5. Build Approval Page
    print("Building Approval Page...")
    build_approval_page_redesigned(doc)
    
    # 6. Build Comprehensive Appendices (Appendix A through E)
    print("Building Comprehensive Appendices (Appendix A through E)...")
    build_appendices_redesigned(doc)
    
    # 7. Save DOCX
    print(f"Saving DOCX to {docx_path}...")
    doc.save(docx_path)
    print(f"DOCX saved successfully! Size: {os.path.getsize(docx_path)} bytes")
    
    # 8. Convert to PDF via Word COM
    print(f"Converting DOCX to PDF via Word COM: {pdf_path}...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wb = word.Documents.Open(docx_path)
        wb.Fields.Update()
        wb.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
        wb.Close()
        print(f"PDF exported successfully! File size: {os.path.getsize(pdf_path)} bytes")
    finally:
        word.Quit()

def get_actual_page_numbers_from_pdf(pdf_path):
    doc_pdf = pymupdf.open(pdf_path)
    romans = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X', 'XI', 'XII']
    bab_pages = {}
    
    for rom in romans:
        for i, page in enumerate(doc_pdf):
            if i < 3:
                continue
            text = page.get_text()
            for line in text.split('\n'):
                line_clean = line.strip()
                if re.match(rf'^[■\u25a0]?\s*BAB\s+{rom}\s*$', line_clean):
                    if rom not in bab_pages:
                        bab_pages[rom] = str(i + 1)
                        break
                        
    approval_page = "30"
    app_a_page = "31"
    app_b_page = "33"
    app_c_page = "34"
    app_d_page = "35"
    app_e_page = "37"
    
    for i, page in enumerate(doc_pdf):
        if i < 3:
            continue
        text = page.get_text()
        for line in text.split('\n'):
            line_clean = line.strip()
            if 'HALAMAN PERSETUJUAN' in line_clean:
                approval_page = str(i + 1)
            elif 'LAMPIRAN A' in line_clean:
                app_a_page = str(i + 1)
            elif 'LAMPIRAN B' in line_clean:
                app_b_page = str(i + 1)
            elif 'LAMPIRAN C' in line_clean:
                app_c_page = str(i + 1)
            elif 'LAMPIRAN D' in line_clean:
                app_d_page = str(i + 1)
            elif 'LAMPIRAN E' in line_clean:
                app_e_page = str(i + 1)
            
    return bab_pages, approval_page, app_a_page, app_b_page, app_c_page, app_d_page, app_e_page

def main():
    docx_path = os.path.abspath("RS_S1_Official_Regulations_Redesigned.docx")
    pdf_path = os.path.abspath("RS_S1_Official_Regulations_Redesigned.pdf")
    cover_img = os.path.abspath("cover_artwork_s1.png")
    
    toc_titles = [
        ("BAB I", "KETENTUAN UMUM"),
        ("BAB II", "FORMAT KOMPETISI, MATCHDAY, DAN KLASEMEN"),
        ("BAB III", "PENGATURAN PERTANDINGAN"),
        ("BAB IV", "SKUAD, REGISTRASI PEMAIN, DAN PERUBAHAN SKUAD"),
        ("BAB V", "PERSYARATAN KARTU TRENDING DAN PEMAIN INDONESIA"),
        ("BAB VI", "PELAKSANAAN, PELAPORAN, DAN PENGESAHAN HASIL"),
        ("BAB VII", "GANGGUAN KONEKSI DAN REKONSTRUKSI PERTANDINGAN"),
        ("BAB VIII", "SPORTIVITAS, BIRDCAGING, DAN PERILAKU TERLARANG"),
        ("BAB IX", "PELANGGARAN, SANKSI, DAN PENYELESAIAN SENGKETA"),
        ("BAB X", "STATISTIK, PENENTUAN JUARA, DAN PENGHARGAAN"),
        ("BAB XI", "WEBSITE, DATA, DAN ADMINISTRASI KOMPETISI"),
        ("BAB XII", "PERUBAHAN REGULASI DAN KETENTUAN PENUTUP"),
    ]
    
    initial_toc_data = [(code, title, "4") for code, title in toc_titles]
    initial_toc_data.append(("HALAMAN PERSETUJUAN", "Tanda Tangan dan Pengesahan Kedua Peserta", "30"))
    initial_toc_data.append(("LAMPIRAN A", "Official Fixture List (MD1–MD20)", "31"))
    initial_toc_data.append(("LAMPIRAN B", "Official Match Report Template & Protocol", "33"))
    initial_toc_data.append(("LAMPIRAN C", "Official League Table & Formula Klasemen", "34"))
    initial_toc_data.append(("LAMPIRAN D", "Disciplinary and Dispute Register", "35"))
    initial_toc_data.append(("LAMPIRAN E", "Regulations Amendment Log & Pre-MD1 Checklist", "37"))
    
    print("=== PASS 1: BUILDING INITIAL DRAFT TO CALIBRATE PAGE NUMBERS ===")
    generate_document(docx_path, pdf_path, cover_img, initial_toc_data)
    
    print("\n=== CALIBRATING EXACT PAGE NUMBERS FROM EXPORTED PDF ===")
    bab_pages, approval_p, app_a_p, app_b_p, app_c_p, app_d_p, app_e_p = get_actual_page_numbers_from_pdf(pdf_path)
    
    final_toc_data = []
    romans = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X', 'XI', 'XII']
    for idx, (code, title) in enumerate(toc_titles):
        rom = romans[idx]
        actual_p = bab_pages.get(rom, "4")
        final_toc_data.append((code, title, actual_p))
        print(f"  {code} -> Page {actual_p}")
        
    final_toc_data.append(("HALAMAN PERSETUJUAN", "Tanda Tangan dan Pengesahan Kedua Peserta", approval_p))
    final_toc_data.append(("LAMPIRAN A", "Official Fixture List (MD1–MD20)", app_a_p))
    final_toc_data.append(("LAMPIRAN B", "Official Match Report Template & Protocol", app_b_p))
    final_toc_data.append(("LAMPIRAN C", "Official League Table & Formula Klasemen", app_c_p))
    final_toc_data.append(("LAMPIRAN D", "Disciplinary and Dispute Register", app_d_p))
    final_toc_data.append(("LAMPIRAN E", "Regulations Amendment Log & Pre-MD1 Checklist", app_e_p))
    
    print(f"  HALAMAN PERSETUJUAN -> Page {approval_p}")
    print(f"  LAMPIRAN A -> Page {app_a_p}")
    print(f"  LAMPIRAN B -> Page {app_b_p}")
    print(f"  LAMPIRAN C -> Page {app_c_p}")
    print(f"  LAMPIRAN D -> Page {app_d_p}")
    print(f"  LAMPIRAN E -> Page {app_e_p}")
    
    print("\n=== PASS 2: GENERATING DEFINITIVE REDESIGNED DOCUMENTS WITH EXACT TOC ===")
    generate_document(docx_path, pdf_path, cover_img, final_toc_data)
    
    print("\n=== VERIFYING FINAL OUTPUTS ===")
    doc_final = pymupdf.open(pdf_path)
    print(f"Final PDF total pages: {len(doc_final)}")
    print(f"Final DOCX file size: {os.path.getsize(docx_path):,} bytes")
    print(f"Final PDF file size: {os.path.getsize(pdf_path):,} bytes")

if __name__ == "__main__":
    main()
