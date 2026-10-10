import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client

# Palette definitions
HEX_PRIMARY_NAVY = "0F1E36"      # Deep Executive Navy
HEX_SECONDARY_NAVY = "1E3A8A"    # Slate Navy
HEX_ACCENT_GOLD = "C5A059"       # Muted Gold
HEX_DARK_TEXT = "1E293B"         # Dark Slate / Charcoal Body Text
HEX_MUTED_TEXT = "64748B"        # Slate Muted Text
HEX_LIGHT_BG = "F8FAFC"          # Light Slate Shading
HEX_ALT_ROW = "F1F5F9"           # Alternating Row Shading
HEX_BORDER = "CBD5E1"            # Subtle Border Gray
HEX_WHITE = "FFFFFF"

COLOR_PRIMARY_NAVY = RGBColor(0x0F, 0x1E, 0x36)
COLOR_SECONDARY_NAVY = RGBColor(0x1E, 0x3A, 0x8A)
COLOR_ACCENT_GOLD = RGBColor(0xC5, 0xA0, 0x59)
COLOR_DARK_TEXT = RGBColor(0x1E, 0x29, 0x3B)
COLOR_MUTED_TEXT = RGBColor(0x64, 0x74, 0x8B)

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=90, bottom=90, left=130, right=130):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, val in borders.items():
        if val:
            sz, b_type, col = val
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{b_type}" w:sz="{sz}" w:space="0" w:color="{col}"/>')
            tcBorders.append(b_xml)
        else:
            b_xml = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b_xml)
    tcPr.append(tcBorders)

def set_table_borders(table, color=HEX_BORDER, sz="4"):
    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

def add_inline_runs(paragraph, text, base_size=Pt(10), base_color=COLOR_DARK_TEXT, base_font="Calibri"):
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = base_font
            run.font.size = base_size
            run.font.bold = True
            run.font.color.rgb = base_color
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = base_font
            run.font.size = base_size
            run.font.italic = True
            run.font.color.rgb = base_color
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size.pt * 0.95)
            run.font.bold = True
            run.font.color.rgb = COLOR_SECONDARY_NAVY
        else:
            run = paragraph.add_run(token)
            run.font.name = base_font
            run.font.size = base_size
            run.font.color.rgb = base_color

def make_callout(doc, title, text, border_color=HEX_ACCENT_GOLD, bg_color=HEX_LIGHT_BG):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_cell_borders(cell, left=(24, "single", border_color), top=None, bottom=None, right=None)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    if title:
        run_title = p.add_run(title + " ")
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(9.5)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_PRIMARY_NAVY
        
    add_inline_runs(p, text, base_size=Pt(9.5), base_color=COLOR_DARK_TEXT)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def format_custom_table(doc, col_widths, headers, rows_data, align_right_cols=[], align_center_cols=[]):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    # Header row
    hdr_cells = table.rows[0].cells
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for i, title in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], HEX_PRIMARY_NAVY)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if i in align_right_cols:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif i in align_center_cols:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data rows
    for r_idx, r_data in enumerate(rows_data):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        bg_col = HEX_ALT_ROW if (r_idx % 2 == 1) else HEX_WHITE
        
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if c_idx in align_right_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif c_idx in align_center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(p, str(val), base_size=Pt(9), base_color=COLOR_DARK_TEXT)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_gold_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="{HEX_ACCENT_GOLD}"/></w:pBdr>')
    pPr.append(pBdr)

def add_header_footer(section):
    section.different_first_page_header_footer = False
    
    # Header
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(4)
    hpPr = hp._p.get_or_add_pPr()
    hpBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="3" w:color="{HEX_BORDER}"/></w:pBdr>')
    hpPr.append(hpBdr)
    
    run_l = hp.add_run("RIVALS SERIES SEASON 1  |  ")
    run_l.font.name = "Calibri"
    run_l.font.size = Pt(8.5)
    run_l.font.bold = True
    run_l.font.color.rgb = COLOR_PRIMARY_NAVY
    
    run_r = hp.add_run("OFFICIAL COMPETITION REGULATIONS")
    run_r.font.name = "Calibri"
    run_r.font.size = Pt(8.5)
    run_r.font.color.rgb = COLOR_MUTED_TEXT

    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(4)
    fp.paragraph_format.space_after = Pt(0)
    fpPr = fp._p.get_or_add_pPr()
    fpBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="3" w:color="{HEX_BORDER}"/></w:pBdr>')
    fpPr.append(fpBdr)
    
    tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="9360"/></w:tabs>')
    fpPr.append(tabs)
    
    run_fl = fp.add_run("Versi 1.0-draft  •  Draft Konsolidasi Resmi")
    run_fl.font.name = "Calibri"
    run_fl.font.size = Pt(8.5)
    run_fl.font.color.rgb = COLOR_MUTED_TEXT
    
    run_tab = fp.add_run("\t")
    
    run_fr = fp.add_run("Halaman ")
    run_fr.font.name = "Calibri"
    run_fr.font.size = Pt(8.5)
    run_fr.font.color.rgb = COLOR_MUTED_TEXT
    
    fld1 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    fp._p.append(fld1)
    
    run_mid = fp.add_run(" dari ")
    run_mid.font.name = "Calibri"
    run_mid.font.size = Pt(8.5)
    run_mid.font.color.rgb = COLOR_MUTED_TEXT
    
    fld2 = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"/>')
    fp._p.append(fld2)

def build_cover_page(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(28)
    p_top.paragraph_format.space_after = Pt(12)
    
    # Monogram Box
    mono_table = doc.add_table(rows=1, cols=1)
    mono_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    mono_table.autofit = False
    mono_cell = mono_table.cell(0, 0)
    mono_cell.width = Inches(2.6)
    set_cell_background(mono_cell, HEX_PRIMARY_NAVY)
    set_cell_margins(mono_cell, top=160, bottom=160, left=180, right=180)
    set_cell_borders(
        mono_cell,
        top=(12, "single", HEX_ACCENT_GOLD),
        bottom=(12, "single", HEX_ACCENT_GOLD),
        left=(12, "single", HEX_ACCENT_GOLD),
        right=(12, "single", HEX_ACCENT_GOLD)
    )
    
    mp1 = mono_cell.paragraphs[0]
    mp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mp1.paragraph_format.space_before = Pt(2)
    mp1.paragraph_format.space_after = Pt(0)
    run_rs = mp1.add_run("RS")
    run_rs.font.name = "Calibri"
    run_rs.font.size = Pt(40)
    run_rs.font.bold = True
    run_rs.font.color.rgb = COLOR_ACCENT_GOLD
    
    mp2 = mono_cell.add_paragraph()
    mp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mp2.paragraph_format.space_before = Pt(0)
    mp2.paragraph_format.space_after = Pt(2)
    run_series = mp2.add_run("RIVALS SERIES")
    run_series.font.name = "Calibri"
    run_series.font.size = Pt(9.5)
    run_series.font.bold = True
    run_series.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Title paragraph
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(28)
    p_title.paragraph_format.space_after = Pt(4)
    run_t1 = p_title.add_run("RIVALS SERIES — SEASON 1")
    run_t1.font.name = "Calibri"
    run_t1.font.size = Pt(24)
    run_t1.font.bold = True
    run_t1.font.color.rgb = COLOR_PRIMARY_NAVY
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    run_t2 = p_sub.add_run("OFFICIAL COMPETITION REGULATIONS")
    run_t2.font.name = "Calibri"
    run_t2.font.size = Pt(16)
    run_t2.font.bold = True
    run_t2.font.color.rgb = COLOR_ACCENT_GOLD
    
    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(18)
    run_desc = p_desc.add_run("Buku Regulasi Resmi Kompetisi Sepak Bola Virtual 1v1")
    run_desc.font.name = "Calibri"
    run_desc.font.size = Pt(11)
    run_desc.font.italic = True
    run_desc.font.color.rgb = COLOR_MUTED_TEXT

    # Metadata card table on cover
    card_table = doc.add_table(rows=6, cols=2)
    card_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    card_table.autofit = False
    set_table_borders(card_table, color=HEX_BORDER, sz="4")
    
    card_data = [
        ("Status Dokumen", "Draft konsolidasi untuk persetujuan kedua peserta"),
        ("Versi Regulasi", "1.0-draft"),
        ("Format Kompetisi", "eFootball™ Dream Team, Friend Match 1v1"),
        ("Jumlah Matchday", "20 Matchday Reguler (MD1–MD20)"),
        ("Peserta Terdaftar", "Peserta A & Peserta B"),
        ("Keberlakuan", "Berlaku penuh setelah persetujuan kedua peserta"),
    ]
    
    for r_idx, (k, v) in enumerate(card_data):
        row = card_table.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(3.6)
        
        set_cell_background(c0, HEX_LIGHT_BG)
        set_cell_background(c1, HEX_WHITE)
        set_cell_margins(c0, top=70, bottom=70, left=120, right=120)
        set_cell_margins(c1, top=70, bottom=70, left=120, right=120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY_NAVY
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK_TEXT
        
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(44)
    p_sp.paragraph_format.space_after = Pt(0)

    p_disc = doc.add_paragraph()
    p_disc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_disc.paragraph_format.space_before = Pt(0)
    p_disc.paragraph_format.space_after = Pt(0)
    p_disc.paragraph_format.line_spacing = 1.15
    
    r_d1 = p_disc.add_run("Kompetisi privat non-komersial yang diselenggarakan secara independen oleh dua peserta.\n")
    r_d1.font.name = "Calibri"
    r_d1.font.size = Pt(8.5)
    r_d1.font.italic = True
    r_d1.font.color.rgb = COLOR_MUTED_TEXT
    
    r_d2 = p_disc.add_run("Dokumen ini tidak berafiliasi secara resmi dengan KONAMI Digital Entertainment, eFootball™, FIFA, UEFA, atau organisasi kompetisi olahraga lainnya.")
    r_d2.font.name = "Calibri"
    r_d2.font.size = Pt(8.5)
    r_d2.font.italic = True
    r_d2.font.color.rgb = COLOR_MUTED_TEXT

def build_doc_control_and_toc(doc):
    # PAGE 2: Kontrol Dokumen
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(6)
    p_h1.paragraph_format.space_after = Pt(2)
    p_h1.paragraph_format.keep_with_next = True
    r = p_h1.add_run("KONTROL DOKUMEN DAN STATUS REGULASI")
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY_NAVY
    
    add_gold_divider(doc)
    
    meta_table = doc.add_table(rows=10, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color=HEX_BORDER, sz="4")
    
    meta_data = [
        ("Nama kompetisi", "Rivals Series Season 1 (RS S1)"),
        ("Status dokumen", "Draft konsolidasi untuk persetujuan kedua peserta"),
        ("Versi dokumen", "1.0-draft"),
        ("Format pertandingan", "eFootball Dream Team, 1v1"),
        ("Jumlah pertandingan reguler", "20 Matchday (MD1–MD20)"),
        ("Pihak peserta", "Peserta A dan Peserta B"),
        ("Competition Administrator", "Ditetapkan oleh kedua peserta"),
        ("Official Website", "Akan ditambahkan setelah website tersedia"),
        ("Tanggal berlaku", "Setelah kedua peserta menyetujui dokumen ini"),
        ("Sumber aturan primer", "rules.md (Konsolidasi mandiri kedua peserta)"),
    ]
    
    for r_idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        
        set_cell_background(c0, HEX_LIGHT_BG)
        set_cell_background(c1, HEX_WHITE)
        set_cell_margins(c0, top=70, bottom=70, left=120, right=120)
        set_cell_margins(c1, top=70, bottom=70, left=120, right=120)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY_NAVY
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK_TEXT
        
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(8)
    sp.paragraph_format.space_after = Pt(2)
    
    make_callout(
        doc,
        title="Prinsip utama:",
        text="Dokumen ini merupakan sumber aturan utama RS S1 berdasarkan rules.md. Buku resmi, website, pengumuman, dan lampiran harus konsisten dengan rules.md. Tidak ada pihak yang boleh mengubah aturan, hasil, atau sanksi secara sepihak."
    )
    
    # PAGE 3: Dedicated DAFTAR ISI page
    doc.add_page_break()
    
    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.space_before = Pt(6)
    p_toc.paragraph_format.space_after = Pt(2)
    p_toc.paragraph_format.keep_with_next = True
    r_toc = p_toc.add_run("DAFTAR ISI")
    r_toc.font.name = "Calibri"
    r_toc.font.size = Pt(14)
    r_toc.font.bold = True
    r_toc.font.color.rgb = COLOR_PRIMARY_NAVY
    
    add_gold_divider(doc)
    
    toc_items = [
        ("BAB I", "Ketentuan Umum"),
        ("BAB II", "Format Kompetisi, Matchday, dan Klasemen"),
        ("BAB III", "Pengaturan Pertandingan"),
        ("BAB IV", "Skuad, Registrasi Pemain, dan Perubahan Skuad"),
        ("BAB V", "Persyaratan Kartu Trending dan Pemain Indonesia"),
        ("BAB VI", "Pelaksanaan, Pelaporan, dan Pengesahan Hasil"),
        ("BAB VII", "Gangguan Koneksi dan Rekonstruksi Pertandingan"),
        ("BAB VIII", "Sportivitas, Birdcaging, dan Perilaku Terlarang"),
        ("BAB IX", "Pelanggaran, Sanksi, dan Penyelesaian Sengketa"),
        ("BAB X", "Statistik, Penentuan Juara, dan Penghargaan"),
        ("BAB XI", "Website, Data, dan Administrasi Kompetisi"),
        ("BAB XII", "Perubahan Regulasi dan Ketentuan Penutup"),
        ("HALAMAN PERSETUJUAN", "Tanda Tangan dan Pengesahan Peserta"),
        ("LAMPIRAN", "Rujukan Appendix A–E (Diterbitkan Terpisah untuk WhatsApp)"),
    ]
    
    toc_table = doc.add_table(rows=len(toc_items), cols=2)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc_table.autofit = False
    set_table_borders(toc_table, color=HEX_BORDER, sz="2")
    
    for idx, (code, title) in enumerate(toc_items):
        row = toc_table.rows[idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        
        bg_col = HEX_LIGHT_BG if (idx % 2 == 1) else HEX_WHITE
        set_cell_background(c0, bg_col)
        set_cell_background(c1, bg_col)
        set_cell_margins(c0, top=55, bottom=55, left=100, right=100)
        set_cell_margins(c1, top=55, bottom=55, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(f"{idx+1}. {code}" if not code.startswith("HALAMAN") and not code.startswith("LAMPIRAN") else f"• {code}")
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY_NAVY
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(f"— {title}")
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK_TEXT
        
        c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def build_approval_page(doc):
    doc.add_page_break()
    
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(10)
    p_h1.paragraph_format.space_after = Pt(2)
    p_h1.paragraph_format.keep_with_next = True
    r = p_h1.add_run("HALAMAN PERSETUJUAN")
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY_NAVY
    
    add_gold_divider(doc)
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(2)
    p_intro.paragraph_format.space_after = Pt(10)
    p_intro.paragraph_format.line_spacing = 1.15
    r_intro = p_intro.add_run("Dengan menyetujui regulasi ini, kedua peserta menyatakan akan mengikuti RS S1 berdasarkan ketentuan yang tercantum dalam dokumen ini.")
    r_intro.font.name = "Calibri"
    r_intro.font.size = Pt(10)
    r_intro.font.color.rgb = COLOR_DARK_TEXT
    
    sig_table = doc.add_table(rows=3, cols=1)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    
    signatories = [
        ("PESERTA A", "Nama: ______________________________\nPersetujuan / Tanda Tangan: ______________________________\nTanggal: ______________________________"),
        ("PESERTA B", "Nama: ______________________________\nPersetujuan / Tanda Tangan: ______________________________\nTanggal: ______________________________"),
        ("COMPETITION ADMINISTRATOR", "Nama: ______________________________\nTanggal Penetapan: ______________________________"),
    ]
    
    for idx, (title, lines_text) in enumerate(signatories):
        row = sig_table.rows[idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        cell = row.cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, HEX_LIGHT_BG)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        set_cell_borders(
            cell,
            left=(20, "single", HEX_PRIMARY_NAVY),
            top=(4, "single", HEX_BORDER),
            bottom=(4, "single", HEX_BORDER),
            right=(4, "single", HEX_BORDER)
        )
        
        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(3)
        r0 = p0.add_run(title)
        r0.font.name = "Calibri"
        r0.font.size = Pt(10)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY_NAVY
        
        for l in lines_text.split('\n'):
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(l)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK_TEXT
            
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

def build_whatsapp_appendix_reference(doc):
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(12)
    p_h2.paragraph_format.space_after = Pt(2)
    p_h2.paragraph_format.keep_with_next = True
    r = p_h2.add_run("RUJUKAN LAMPIRAN ADMINISTRATIF (APPENDIX A–E)")
    r.font.name = "Calibri"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY_NAVY
    
    add_gold_divider(doc)
    
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(2)
    p_note.paragraph_format.space_after = Pt(4)
    p_note.paragraph_format.line_spacing = 1.15
    add_inline_runs(
        p_note,
        "Lampiran administratif **Appendix A sampai Appendix E** diterbitkan dan didistribusikan secara terpisah melalui dokumen **`RS_S1_Appendices_WhatsApp.md`** agar praktis disalin dan digunakan langsung dalam percakapan administrasi WhatsApp antarpeserta. Seluruh lampiran tersebut berstatus sebagai perangkat operasional subordinat dan **tidak boleh bertentangan dengan ketentuan buku regulasi ini**.",
        base_size=Pt(9.5)
    )
    
    appendices_summary = [
        ("Appendix A — Official Fixture List", "Template pemantauan jadwal dan status 20 Matchday (MD1–MD20) tanpa tanggal atau hasil rekaan."),
        ("Appendix B — Official Match Report", "Format pelaporan hasil pertandingan pasca-laga, konfirmasi kedua pihak, catatan rekonstruksi gangguan, dan bukti DSP."),
        ("Appendix C — Official League Table", "Tautan website resmi, panduan kolom klasemen (MP, W, D, L, GF, GA, GD, PTS, ID Credits, Sanctions), dan formula poin."),
        ("Appendix D — Disciplinary and Dispute Register", "Format pelaporan sengketa, bukti gameplay, batas tanggapan 48 jam, adjudikasi AI, dan penegakan aturan status quo."),
        ("Appendix E — Regulations Amendment Log & Pre-MD1 Checklist", "Log pencatatan amendemen resmi regulasi serta checklist penetapan pra-kompetisi (AI Adjudicator, alat undian digital, dan snapshot skuad)."),
    ]
    
    for code, desc in appendices_summary:
        p_app = doc.add_paragraph()
        p_app.paragraph_format.left_indent = Inches(0.25)
        p_app.paragraph_format.first_line_indent = Inches(-0.25)
        p_app.paragraph_format.space_before = Pt(1.5)
        p_app.paragraph_format.space_after = Pt(1.5)
        p_app.paragraph_format.line_spacing = 1.15
        
        r_b = p_app.add_run(f"• {code}: ")
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_SECONDARY_NAVY
        
        r_d = p_app.add_run(desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9)
        r_d.font.color.rgb = COLOR_DARK_TEXT

def parse_and_render_body(doc, body_lines):
    idx = 0
    total = len(body_lines)
    is_first_bab = True
    
    while idx < total:
        line = body_lines[idx].rstrip('\r\n')
        s = line.strip()
        
        if not s:
            idx += 1
            continue
            
        # 1. Heading 1: BAB
        if s.startswith('# BAB '):
            if is_first_bab:
                # Page break before BAB I (starts on fresh page after TOC)
                doc.add_page_break()
                is_first_bab = False
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(22)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(s.lstrip('#').strip())
            run.font.name = "Calibri"
            run.font.size = Pt(13.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY_NAVY
            
            add_gold_divider(doc)
            idx += 1
            continue
            
        # 2. Heading 2: Pasal
        if s.startswith('## Pasal '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(s.lstrip('#').strip())
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY_NAVY
            idx += 1
            continue
            
        # 3. Heading 3: Ayat / Sub-bagian
        if s.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            
            run = p.add_run(s.lstrip('#').strip())
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = COLOR_SECONDARY_NAVY
            idx += 1
            continue
            
        # 4. Divider '---'
        if s.startswith('---'):
            idx += 1
            continue
            
        # 5. Callout block
        if s.startswith('> '):
            callout_text = s[2:].strip()
            title = ""
            m = re.match(r'^\*\*(.*?)\*\*\s*(.*)$', callout_text)
            if m:
                title = m.group(1)
                callout_text = m.group(2)
            make_callout(doc, title=title, text=callout_text)
            idx += 1
            continue
            
        # 6. Table block
        if s.startswith('|'):
            table_lines = []
            while idx < total and body_lines[idx].strip().startswith('|'):
                table_lines.append(body_lines[idx].strip())
                idx += 1
                
            hdr_line = table_lines[0]
            headers = [c.strip() for c in hdr_line.strip('|').split('|')]
            
            align_line = table_lines[1]
            align_specs = [c.strip() for c in align_line.strip('|').split('|')]
            align_right = [i for i, spec in enumerate(align_specs) if spec.endswith(':') and not spec.startswith(':')]
            align_center = [i for i, spec in enumerate(align_specs) if spec.startswith(':') and spec.endswith(':')]
            
            rows_data = []
            for tl in table_lines[2:]:
                r_cells = [c.strip() for c in tl.strip('|').split('|')]
                rows_data.append(r_cells)
                
            num_cols = len(headers)
            if num_cols == 2:
                if headers == ["Hasil", "Poin peserta"]:
                    col_widths = [3.5, 3.0]
                elif headers == ["Kolom", "Arti"]:
                    col_widths = [1.8, 4.7]
                elif headers == ["Pengaturan", "Ketentuan"]:
                    col_widths = [3.2, 3.3]
                elif headers == ["Pengaturan", "Status"]:
                    col_widths = [3.5, 3.0]
                elif headers == ["Pemakaian pemain Indonesia", "Kredit"]:
                    col_widths = [5.0, 1.5]
                else:
                    col_widths = [3.25, 3.25]
            elif num_cols == 3:
                col_widths = [2.2, 2.2, 2.1]
            else:
                w = round(6.5 / num_cols, 2)
                col_widths = [w] * num_cols
                col_widths[-1] = 6.5 - sum(col_widths[:-1])
                
            format_custom_table(
                doc,
                col_widths=col_widths,
                headers=headers,
                rows_data=rows_data,
                align_right_cols=align_right,
                align_center_cols=align_center
            )
            continue
            
        # 7. Lettered sub-items: '   a. ' or 'a. '
        m_let = re.match(r'^\s*([a-z])\.\s+(.*)$', s)
        if m_let:
            letter = m_let.group(1)
            content = m_let.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.45)
            p.paragraph_format.first_line_indent = Inches(-0.20)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2.0)
            p.paragraph_format.line_spacing = 1.15
            
            run_lbl = p.add_run(f"{letter}. ")
            run_lbl.font.name = "Calibri"
            run_lbl.font.size = Pt(9.5)
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = COLOR_SECONDARY_NAVY
            
            add_inline_runs(p, content, base_size=Pt(9.5))
            idx += 1
            continue
            
        # 8. Numbered items: '1. '
        m_num = re.match(r'^(\d+)\.\s+(.*)$', s)
        if m_num:
            num = m_num.group(1)
            content = m_num.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2.0)
            p.paragraph_format.line_spacing = 1.15
            
            run_lbl = p.add_run(f"{num}. ")
            run_lbl.font.name = "Calibri"
            run_lbl.font.size = Pt(9.5)
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = COLOR_PRIMARY_NAVY
            
            add_inline_runs(p, content, base_size=Pt(9.5))
            idx += 1
            continue
            
        # 9. Bullet items: '- ' or '* '
        if s.startswith('- ') or s.startswith('* '):
            content = s[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2.0)
            p.paragraph_format.line_spacing = 1.15
            
            run_lbl = p.add_run("• ")
            run_lbl.font.name = "Calibri"
            run_lbl.font.size = Pt(9.5)
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = COLOR_SECONDARY_NAVY
            
            add_inline_runs(p, content, base_size=Pt(9.5))
            idx += 1
            continue
            
        # 10. Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.15
        add_inline_runs(p, s, base_size=Pt(9.5))
        idx += 1

def main():
    print("Reading rules.md...")
    with open('rules.md', 'r', encoding='utf-8') as f:
        full_text = f.read()
        
    lines = full_text.splitlines()
    
    bab1_idx = -1
    approval_idx = -1
    
    for i, l in enumerate(lines):
        if l.strip().startswith('# BAB I '):
            bab1_idx = i
            break
            
    for i, l in enumerate(lines):
        if l.strip().startswith('# HALAMAN PERSETUJUAN'):
            approval_idx = i
            break
            
    print(f"BAB I starts at line {bab1_idx + 1}")
    print(f"HALAMAN PERSETUJUAN starts at line {approval_idx + 1}")
    
    body_lines = lines[bab1_idx:approval_idx]
    
    doc = docx.Document()
    
    # 1. Build Cover Page
    print("Building Cover Page...")
    build_cover_page(doc)
    
    # 2. Add Section Break for Body
    sec_body = doc.add_section(WD_SECTION.NEW_PAGE)
    sec_body.top_margin = Inches(1.0)
    sec_body.bottom_margin = Inches(1.0)
    sec_body.left_margin = Inches(1.0)
    sec_body.right_margin = Inches(1.0)
    
    add_header_footer(sec_body)
    
    # 3. Build Document Control & TOC
    print("Building Document Control & TOC...")
    build_doc_control_and_toc(doc)
    
    # 4. Parse & Render BAB I to BAB XII
    print("Parsing & Rendering BAB I through BAB XII...")
    parse_and_render_body(doc, body_lines)
    
    # 5. Build Approval Page
    print("Building Approval Page...")
    build_approval_page(doc)
    
    # 6. Build WhatsApp Appendix Reference
    print("Building WhatsApp Appendix Reference...")
    build_whatsapp_appendix_reference(doc)
    
    # Save DOCX
    docx_path = os.path.abspath("RS_S1_Official_Regulations.docx")
    print(f"Saving DOCX to {docx_path}...")
    doc.save(docx_path)
    print("DOCX saved successfully!")
    
    # Convert to PDF via Word COM
    pdf_path = os.path.abspath("RS_S1_Official_Regulations.pdf")
    print(f"Converting DOCX to PDF via Word COM: {pdf_path}...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wb = word.Documents.Open(docx_path)
        wb.Fields.Update()
        wb.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
        wb.Close()
        print(f"PDF exported successfully! File size: {os.path.getsize(pdf_path)} bytes")
    finally:
        word.Quit()

if __name__ == "__main__":
    main()
