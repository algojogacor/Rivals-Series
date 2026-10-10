import docx
from docx.shared import Mm, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client
from PIL import Image
import pymupdf

doc = docx.Document()
sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = Mm(0)
sec.bottom_margin = Mm(0)
sec.left_margin = Mm(0)
sec.right_margin = Mm(0)
sec.header_distance = Mm(0)
sec.footer_distance = Mm(0)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = Pt(1)
r = p.add_run()
inline_shape = r.add_picture('cover_artwork_s1.png', width=Mm(210), height=Mm(297))

drawing = r._r.xpath('.//w:drawing')[0]
inline = drawing.xpath('.//wp:inline')[0]

extent = inline.xpath('.//wp:extent')[0]
docPr = inline.xpath('.//wp:docPr')[0]
cNvGraphicFramePr = inline.xpath('.//wp:cNvGraphicFramePr')[0]
graphic = inline.xpath('.//a:graphic')[0]

anchor_xml = (
    f'<wp:anchor {nsdecls("wp")} {nsdecls("a")} distT="0" distB="0" distL="0" distR="0" '
    f'simplePos="0" relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
    f'<wp:simplePos x="0" y="0"/>'
    f'<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    f'</wp:anchor>'
)
anchor = parse_xml(anchor_xml)
anchor.append(extent)
anchor.append(parse_xml(f'<wp:wrapNone {nsdecls("wp")}/>'))
anchor.append(docPr)
anchor.append(cNvGraphicFramePr)
anchor.append(graphic)

drawing.remove(inline)
drawing.append(anchor)

docx_path = r'D:\Projects\Rivals-Series\test_cover_anchor.docx'
pdf_path = r'D:\Projects\Rivals-Series\test_cover_anchor.pdf'
doc.save(docx_path)

word = win32com.client.Dispatch('Word.Application')
word.Visible = False
wb = word.Documents.Open(docx_path)
wb.SaveAs(pdf_path, FileFormat=17)
wb.Close()
word.Quit()

doc_pdf = pymupdf.open(pdf_path)
print('PDF pages:', len(doc_pdf))
pix = doc_pdf[0].get_pixmap(dpi=150)
pix.save('test_cover_anchor.png')

im = Image.open('test_cover_anchor.png')
for y in range(im.height):
    pix_val = im.getpixel((im.width // 2, y))
    if pix_val != (255, 255, 255):
        print(f'First non-white pixel at y = {y}: {pix_val}')
        break
